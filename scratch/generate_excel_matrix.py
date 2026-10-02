import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_access_matrix_excel(filepath):
    wb = openpyxl.Workbook()

    # Define Styles
    header_fill = PatternFill(start_color="3730A3", end_color="3730A3", fill_type="solid") # Indigo 800
    section_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid") # Indigo 50
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Calibri", size=10, color="334155")
    mono_font = Font(name="Consolas", size=10, color="312E81")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="3730A3"))

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # =========================================================================
    # ABA 1: Matriz de Usuários e Acessos
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Matriz de Acessos"
    ws1.views.sheetView[0].showGridLines = True

    # Title block
    ws1.merge_cells("A1:O1")
    ws1["A1"] = "SICOU - MATRIZ CORPORATIVA DE ACESSOS E PERFIS DE USUÁRIO"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A2:O2")
    ws1["A2"] = "Ambiente de Testes | Empresa: Cooperativa Aurora Central (CNPJ: 12.345.678/0001-90)"
    ws1["A2"].font = subtitle_font
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers = [
        "Nível / Papel",
        "Nome Completo",
        "E-mail de Login",
        "Senha",
        "Role Sistema",
        "Empresa",
        "Unidade",
        "Área da Sede",
        "Ver Área",
        "Gerenciar Área",
        "Gerenciar Workflows",
        "Tramitar Processos",
        "Publicar Informativos",
        "Gerenciar Guia",
        "Descrição de Poderes e Escopo no Sistema"
    ]

    ws1.row_dimensions[4].height = 28
    for col_num, header in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    data_users = [
        [
            "Super Admin", "Super Administrador", "admin@sicou.com", "Admin123", "SUPER_ADMIN",
            "Global (Todas)", "Todas", "Todas",
            "SIM", "SIM", "SIM", "SIM", "SIM", "SIM",
            "Acesso irrestrito a todas as empresas, áreas, unidades, controle de acessos e criação de workflows globais."
        ],
        [
            "Admin da Empresa", "Administrador Cooperativa Aurora", "admin.aurora@sicou.com", "Senha@123", "COMPANY_ADMIN",
            "Cooperativa Aurora", "Sede / Todas", "Todas as Áreas",
            "SIM", "SIM", "SIM", "SIM", "SIM", "SIM",
            "Gestão completa de usuários, áreas, unidades e workflows de todas as áreas da Cooperativa Aurora."
        ],
        # Jurídico
        [
            "Gestor Jurídico", "Gestor Jurídico", "gestor.juridico@aurora.com", "Senha@123", "HEADQUARTER_USER",
            "Cooperativa Aurora", "Sede", "Jurídico & Compliance",
            "SIM", "SIM", "SIM", "SIM", "SIM", "SIM",
            "Gestão total da área Jurídica: cria campos, locais de processo, árvores, homologa, clona versões e tramita chamados."
        ],
        [
            "Analista Jurídico", "Analista Jurídico", "analista.juridico@aurora.com", "Senha@123", "HEADQUARTER_USER",
            "Cooperativa Aurora", "Sede", "Jurídico & Compliance",
            "SIM", "NÃO", "NÃO", "NÃO", "NÃO", "NÃO",
            "Apenas visualização na área Jurídica: consulta processos e catálogos sem permissão para criar, editar ou tramitar."
        ],
        # RH
        [
            "Gestor RH", "Gestor Recursos Humanos", "gestor.rh@aurora.com", "Senha@123", "HEADQUARTER_USER",
            "Cooperativa Aurora", "Sede", "Recursos Humanos",
            "SIM", "SIM", "SIM", "SIM", "SIM", "SIM",
            "Gestão total da área de RH: cria campos, locais, árvores de RH e tramita chamados de pessoas e DP."
        ],
        [
            "Analista RH", "Analista Recursos Humanos", "analista.rh@aurora.com", "Senha@123", "HEADQUARTER_USER",
            "Cooperativa Aurora", "Sede", "Recursos Humanos",
            "SIM", "NÃO", "NÃO", "NÃO", "NÃO", "NÃO",
            "Apenas visualização na área de RH: consulta dados sem poder de gestão."
        ],
        # Financeiro
        [
            "Gestor Financeiro", "Gestor Financeiro", "gestor.financeiro@aurora.com", "Senha@123", "HEADQUARTER_USER",
            "Cooperativa Aurora", "Sede", "Financeiro & Contábil",
            "SIM", "SIM", "SIM", "SIM", "SIM", "SIM",
            "Gestão total da área Financeira: cria campos, locais, árvores financeiras e atende solicitações da controladoria."
        ],
        [
            "Analista Financeiro", "Analista Financeiro", "analista.financeiro@aurora.com", "Senha@123", "HEADQUARTER_USER",
            "Cooperativa Aurora", "Sede", "Financeiro & Contábil",
            "SIM", "NÃO", "NÃO", "NÃO", "NÃO", "NÃO",
            "Apenas visualização na área Financeira: consulta dados sem poder de alteração."
        ],
        # Unidades
        [
            "Gestor Cascavel", "Gestor Unidade Cascavel", "gestor.cascavel@aurora.com", "Senha@123", "UNIT_USER",
            "Cooperativa Aurora", "Unidade 01 - Cascavel", "-",
            "-", "-", "-", "-", "-", "-",
            "Colaborador da Unidade Cascavel: abre processos para as áreas da sede e acompanha na aba Minhas Solicitações."
        ],
        [
            "Atendente Cascavel", "Atendente Unidade Cascavel", "atendente.cascavel@aurora.com", "Senha@123", "UNIT_USER",
            "Cooperativa Aurora", "Unidade 01 - Cascavel", "-",
            "-", "-", "-", "-", "-", "-",
            "Atendente da Unidade Cascavel: abertura de solicitações e consulta dos seus próprios chamados abertos."
        ],
        [
            "Gestor Toledo", "Gestor Unidade Toledo", "gestor.toledo@aurora.com", "Senha@123", "UNIT_USER",
            "Cooperativa Aurora", "Unidade 02 - Toledo", "-",
            "-", "-", "-", "-", "-", "-",
            "Colaborador da Unidade Toledo: abertura de processos vinculados à Unidade Toledo."
        ],
        [
            "Atendente Toledo", "Atendente Unidade Toledo", "atendente.toledo@aurora.com", "Senha@123", "UNIT_USER",
            "Cooperativa Aurora", "Unidade 02 - Toledo", "-",
            "-", "-", "-", "-", "-", "-",
            "Atendente da Unidade Toledo: abertura de chamados e acompanhamento da Unidade Toledo."
        ],
        [
            "Gestor Foz", "Gestor Unidade Foz", "gestor.foz@aurora.com", "Senha@123", "UNIT_USER",
            "Cooperativa Aurora", "Unidade 03 - Foz do Iguaçu", "-",
            "-", "-", "-", "-", "-", "-",
            "Colaborador da Unidade Foz do Iguaçu: abertura de solicitações e acompanhamento."
        ],
        [
            "Atendente Foz", "Atendente Unidade Foz", "atendente.foz@aurora.com", "Senha@123", "UNIT_USER",
            "Cooperativa Aurora", "Unidade 03 - Foz do Iguaçu", "-",
            "-", "-", "-", "-", "-", "-",
            "Atendente da Unidade Foz do Iguaçu: abertura de chamados e acompanhamento."
        ],
    ]

    for row_idx, row_data in enumerate(data_users, 5):
        ws1.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)
        row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if is_even else None

        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            # Alignments & Fonts
            if col_idx in [1]:
                cell.font = bold_font
                cell.alignment = left_align
            elif col_idx in [3, 4, 5]:
                cell.font = mono_font
                cell.alignment = left_align
            elif col_idx in [9, 10, 11, 12, 13, 14]:
                cell.alignment = center_align
                if val == "SIM":
                    cell.font = Font(name="Calibri", size=10, bold=True, color="15803D") # Green
                elif val == "NÃO":
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C") # Red
                else:
                    cell.font = regular_font
            elif col_idx == 15:
                cell.font = regular_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.font = regular_font
                cell.alignment = left_align

    # Auto-adjust column widths for Sheet 1
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        if col[0].column == 15:
            ws1.column_dimensions[col_letter].width = 45
        elif col[0].column == 2:
            ws1.column_dimensions[col_letter].width = 30
        elif col[0].column in [3, 4]:
            ws1.column_dimensions[col_letter].width = 28
        else:
            ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # =========================================================================
    # ABA 2: Estrutura Organizacional
    # =========================================================================
    ws2 = wb.create_sheet(title="Estrutura da Empresa")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:D1")
    ws2["A1"] = "ESTRUTURA CADASTRADA - COOPERATIVA AURORA CENTRAL"
    ws2["A1"].font = title_font

    # Unidades table
    ws2["A3"] = "1. UNIDADES / FILIAIS"
    ws2["A3"].font = bold_font

    unit_headers = ["Nome da Unidade", "Código", "Documento (CNPJ)", "Tipo"]
    for i, h in enumerate(unit_headers, 1):
        c = ws2.cell(row=4, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border

    units_data = [
        ["Unidade 01 - Cascavel", "UNID-CSC", "12.345.678/0002-71", "Filial Operacional"],
        ["Unidade 02 - Toledo", "UNID-TLD", "12.345.678/0003-52", "Filial Operacional"],
        ["Unidade 03 - Foz do Iguaçu", "UNID-FOZ", "12.345.678/0004-33", "Filial Operacional"],
    ]
    for r_idx, r_data in enumerate(units_data, 5):
        for c_idx, val in enumerate(r_data, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=val)
            c.font = regular_font
            c.border = thin_border

    # Áreas table
    ws2["A10"] = "2. ÁREAS DA SEDE (SEM MÓDULOS ATRIBUÍDOS)"
    ws2["A10"].font = bold_font

    area_headers = ["Nome da Área", "Slug / Identificador", "Finalidade", "Módulos Habilitados"]
    for i, h in enumerate(area_headers, 1):
        c = ws2.cell(row=11, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border

    areas_data = [
        ["Jurídico & Compliance", "juridico-compliance", "Gestão jurídica, contratos e regulatório", "Nenhum"],
        ["Recursos Humanos", "recursos-humanos", "Gestão de pessoas, DP e cultura", "Nenhum"],
        ["Financeiro & Contábil", "financeiro-contabil", "Controladoria, contas e operações", "Nenhum"],
    ]
    for r_idx, r_data in enumerate(areas_data, 12):
        for c_idx, val in enumerate(r_data, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=val)
            c.font = regular_font
            c.border = thin_border

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 25)

    # Save
    wb.save(filepath)
    print(f"Planilha gerada com sucesso em: {filepath}")

if __name__ == "__main__":
    create_access_matrix_excel("c:\\Users\\kinha\\sicou\\matriz_de_acessos.xlsx")
